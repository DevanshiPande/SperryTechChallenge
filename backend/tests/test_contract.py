"""Every file in contract/examples/ must match the real response structurally: every key in the example exists in the
response (recursively) with a compatible type. Additive keys in the response are allowed."""
import json
import os
from pathlib import Path

import pytest

EXAMPLES = Path(__file__).resolve().parents[1] / "contract" / "examples"
OV_TOP = "OVL_DESC_3__GPC_2"  # the fixture's #1 opportunity (the mock's OVL_3)
OV_CROSS = "OVL_DESC_2__GPC_1"  # the shared-Thurmond crossing (the mock's OVL_2)

WHATIF_PINS = {"name": "Proposed: Hardeeville 230 kV Tap", "utility": "DESC", "voltage_kv": 230,
               "endpoints": [{"name": "Hardeeville", "lat": 32.2866, "lon": -81.0801}],
               "construction_window": {"start": "2027-09-01", "end": "2029-03-01"}, "max_shift_months": 24}
WHATIF_TEXT = {"name": "Corridor upgrade", "location_text": "Near Hardeeville, SC",
               "construction_window": {"start": "2027-06-01", "end": "2027-12-15"}}
DRAFT_PROJECT = {"kind": "project", "draft": {},
                 "text": "We're planning a corridor upgrade near Hardeeville from June to December 2027 with a nightly "
                         "lane closure on US-17 northbound, 9 PM to 5 AM"}

CASES = {
    "GET_health.json": ("GET", "/health", None, 200),
    "GET_meta.json": ("GET", "/meta", None, 200),
    "GET_projects.json": ("GET", "/projects", None, 200),
    "GET_projects_filtered.json": ("GET", "/projects?utility=GPC&year_from=2025", None, 200),
    "GET_project_DESC_3.json": ("GET", "/projects/DESC_3", None, 200),
    "GET_overlaps.json": ("GET", "/overlaps", None, 200),
    "GET_overlaps_filtered.json": ("GET", "/overlaps?max_mi=10&sort=distance", None, 200),
    "GET_overlap_OVL_2.json": ("GET", f"/overlaps/{OV_CROSS}", None, 200),
    "GET_overlap_OVL_3.json": ("GET", f"/overlaps/{OV_TOP}", None, 200),
    "GET_quality.json": ("GET", "/quality", None, 200),
    "POST_whatif.json": ("POST", "/whatif", WHATIF_PINS, 200),
    "POST_whatif_location.json": ("POST", "/whatif", WHATIF_TEXT, 200),
    "POST_brief.json": ("POST", f"/overlaps/{OV_TOP}/brief", {"mode": "summary"}, 200),
    "POST_brief_plan.json": ("POST", f"/overlaps/{OV_TOP}/brief", {"mode": "plan"}, 200),
    "POST_brief_risks.json": ("POST", f"/overlaps/{OV_TOP}/brief", {"mode": "risks"}, 200),
    "POST_draft_project.json": ("POST", "/draft", DRAFT_PROJECT, 200),
    "POST_draft_resource.json": ("POST", "/draft", {"kind": "resource", "draft": {},
                                                    "text": "Two bucket trucks near Savannah June 3-10 2027 at $1,250 "
                                                            "each"}, 200),
    "POST_draft_job.json": ("POST", "/draft", {"kind": "job", "draft": {},
                                               "text": "Need 3 certified lineworkers near Hardeeville from June to "
                                                       "December 2027"}, 200),
    "POST_scenario_assist.json": ("POST", "/scenario/assist",
                                  {"overlap_id": OV_TOP, "message": "Try sharing 2 bucket trucks for 5 days and a "
                                                                    "single laydown yard"}, 200),
    "POST_ask.json": ("POST", "/ask", {"question": "How much could Jasper and McIntosh save by sharing costs?"}, 200),
    "ERROR_400.json": ("GET", "/overlaps?sort=bad", None, 400),
    "ERROR_400_draft.json": ("POST", "/draft", {"kind": "car", "text": "x"}, 400),
    "ERROR_400_whatif.json": ("POST", "/whatif", {}, 400),
    "ERROR_404.json": ("GET", "/projects/GPC_99", None, 404),
}


def same_shape(example, actual, path="$"):
    """Returns a list of mismatches."""
    if example is None or actual is None:
        return []
    if isinstance(example, bool) or isinstance(actual, bool):
        return [] if isinstance(example, bool) == isinstance(actual, bool) else [f"{path}: bool vs {type(actual)}"]
    if isinstance(example, (int, float)):
        return [] if isinstance(actual, (int, float)) else [f"{path}: number vs {type(actual).__name__}"]
    if isinstance(example, str):
        return [] if isinstance(actual, str) else [f"{path}: string vs {type(actual).__name__}"]
    if isinstance(example, list):
        if not isinstance(actual, list):
            return [f"{path}: list vs {type(actual).__name__}"]
        if example and actual:
            return same_shape(example[0], actual[0], path + "[0]")
        return []
    if isinstance(example, dict):
        if not isinstance(actual, dict):
            return [f"{path}: object vs {type(actual).__name__}"]
        out = []
        for k, v in example.items():
            if k not in actual:
                out.append(f"{path}.{k}: missing")
            elif path.endswith(".geometry") and k == "coordinates":
                continue  # Point vs LineString nesting; both allowed by the contract
            else:
                out.extend(same_shape(v, actual[k], f"{path}.{k}"))
        return out
    return []


def _call(client, method, url, body):
    return client.get(url) if method == "GET" else client.post(url, json=body)


ORIGINAL = set(CASES) | {"ERROR_503_ask.json", "POST_brief_fallback.json"}  # the frontend's contract


def test_every_original_example_is_covered():
    assert ORIGINAL <= {p.name for p in EXAMPLES.glob("*.json")}


@pytest.mark.parametrize("name", sorted(CASES))
def test_example_shape(client_ro, name):
    method, url, body, status = CASES[name]
    r = _call(client_ro, method, url, body)
    assert r.status_code == status, r.text
    example = json.loads((EXAMPLES / name).read_text(encoding="utf-8"))
    assert same_shape(example, r.json()) == []


def test_every_list_item_has_the_shape(client_ro):
    for name, url in (("GET_projects.json", "/projects"), ("GET_overlaps.json", "/overlaps")):
        example = json.loads((EXAMPLES / name).read_text(encoding="utf-8"))[0]
        for item in client_ro.get(url).json():
            assert same_shape(example, item) == [], item["id"]


def test_brief_fallback_shape(client_ro, monkeypatch):
    monkeypatch.setenv("FAKE_GEMINI", "0")  # no key -> template fallback
    r = client_ro.post(f"/overlaps/{OV_CROSS}/brief", json={"mode": "summary"})
    assert r.status_code == 200
    body = r.json()
    assert body["source"] == "fallback"
    example = json.loads((EXAMPLES / "POST_brief_fallback.json").read_text(encoding="utf-8"))
    assert same_shape(example, body) == []


def test_ask_503_shape(client_ro, monkeypatch):
    monkeypatch.setenv("FAKE_GEMINI", "0")
    r = client_ro.post("/ask", json={"question": "anything"})
    assert r.status_code == 503
    example = json.loads((EXAMPLES / "ERROR_503_ask.json").read_text(encoding="utf-8"))
    assert same_shape(example, r.json()) == []
    assert r.json()["error"]["code"] == "GEMINI_UNAVAILABLE"


def test_export_matches_sperry_template(client_ro):
    import io

    import openpyxl
    r = client_ro.get("/export/overlaps.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    template = openpyxl.load_workbook(Path(__file__).resolve().parents[1] / "data" / "raw" / "Projects_Overlaps.xlsx")
    for sheet in ("overlaps", "projects"):
        head = [c.value for c in next(wb[sheet].iter_rows(max_row=1))]
        want = [c.value for c in next(template[sheet].iter_rows(max_row=1))]
        assert head == want, sheet
    rows = list(wb["overlaps"].iter_rows(min_row=2, values_only=True))
    assert [(r[4], r[7], r[1], r[2]) for r in rows] == [
        ("DESC_2", "GPC_1", 4.09, 3074), ("DESC_3", "GPC_2", 5.65, 152), ("DESC_3", "GPC_3", 7.55, 517),
        ("DESC_1", "GPC_1", 8.01, 3074), ("DESC_5", "GPC_2", 14.34, 365), ("DESC_5", "GPC_3", 14.81, 730)]
    projects = {r[0]: r for r in wb["projects"].iter_rows(min_row=2, values_only=True)}
    assert projects["DESC_3"][13] == 2 and set(projects["DESC_3"][14:16]) == {"GPC_2", "GPC_3"}
